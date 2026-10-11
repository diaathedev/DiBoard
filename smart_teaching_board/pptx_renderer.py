"""
PowerPoint (.pptx) Handler for Smart Teaching Board.
"""

import os
from typing import Optional
from PySide6.QtGui import QPixmap, QImage, QPainter, QColor, QFont, QPen
from PySide6.QtCore import QRectF, Qt

try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False


class PPTXRenderer:
    def __init__(self, file_path: str):
        self.file_path = file_path
        self.prs = None
        self.slide_count = 0

        if not PPTX_AVAILABLE:
            raise RuntimeError("python-pptx is not installed. Please run: pip install python-pptx")

        self.prs = Presentation(file_path)
        self.slide_count = len(self.prs.slides)
        self.slide_width = int(self.prs.slide_width.inches * 100)
        self.slide_height = int(self.prs.slide_height.inches * 100)

    def render_slide_to_pixmap(self, slide_index: int) -> Optional[QPixmap]:
        if slide_index < 0 or slide_index >= self.slide_count:
            return None

        slide = self.prs.slides[slide_index]
        w = max(800, self.slide_width)
        h = max(600, self.slide_height)

        img = QImage(w, h, QImage.Format_ARGB32)
        img.fill(QColor(255, 255, 255))

        painter = QPainter(img)
        painter.setRenderHint(QPainter.Antialiasing, True)
        painter.setRenderHint(QPainter.TextAntialiasing, True)

        painter.setPen(QPen(QColor(226, 232, 240), 1))
        painter.drawRect(0, 0, w - 1, h - 1)

        painter.setPen(QColor(148, 163, 184))
        painter.setFont(QFont("Arial", 12))
        painter.drawText(w - 120, h - 30, f"Slide {slide_index + 1}/{self.slide_count}")

        for shape in slide.shapes:
            if shape.has_text_frame:
                left = int(shape.left.inches * 100)
                top = int(shape.top.inches * 100)
                width = int(shape.width.inches * 100)
                height = int(shape.height.inches * 100)

                rect = QRectF(left, top, width, height)
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        painter.setPen(QColor(30, 41, 59))
                        painter.setFont(QFont("Arial", 18 if top < 120 else 14))
                        painter.drawText(rect, Qt.TextWordWrap, text)
                        rect.setTop(rect.top() + 28)

        painter.end()
        return QPixmap.fromImage(img)
